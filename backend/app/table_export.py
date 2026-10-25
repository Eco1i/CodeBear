from __future__ import annotations

import csv
import json
import re
from io import BytesIO, StringIO
from typing import Any, Literal, cast

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill

TableExportFormat = Literal["xlsx", "csv", "json", "md"]

_LABELS: dict[str, dict[str, Any]] = {
    "zh-CN": {
        "sheet": "表结构",
        "table_name": "表名称",
        "table_code": "表代码",
        "table_description": "表描述",
        "ordinal": "序号",
        "primary_key": "主键",
        "field_code": "字段英文名",
        "field_description": "字段描述",
        "data_type": "数据类型",
        "length": "长度",
        "nullable": "可空",
        "default_value": "缺省值",
        "comment": "字段备注",
        "yes": "是",
        "no": "否",
    },
    "en-US": {
        "sheet": "Table Structure",
        "table_name": "Table Name",
        "table_code": "Table Code",
        "table_description": "Table Description",
        "ordinal": "No.",
        "primary_key": "Primary Key",
        "field_code": "Field Name",
        "field_description": "Description",
        "data_type": "Data Type",
        "length": "Length",
        "nullable": "Nullable",
        "default_value": "Default Value",
        "comment": "Comment",
        "yes": "Yes",
        "no": "No",
    },
}

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def table_export_stem(detail: dict[str, Any]) -> str:
    """Return a filesystem-safe, cross-platform name for a table export."""
    raw_name = str(detail.get("code") or detail.get("name") or "table").strip()
    safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", raw_name)
    safe_name = safe_name[:120].strip(" .")
    if not safe_name:
        safe_name = "table"
    if safe_name.split(".", maxsplit=1)[0].upper() in {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        *(f"COM{index}" for index in range(1, 10)),
        *(f"LPT{index}" for index in range(1, 10)),
    }:
        safe_name = f"_{safe_name}"
    return safe_name


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _fields(detail: dict[str, Any]) -> list[dict[str, Any]]:
    fields = detail.get("fields", [])
    if not isinstance(fields, list):
        return []
    return sorted(
        (field for field in fields if isinstance(field, dict)),
        key=_ordinal,
    )


def _ordinal(field: dict[str, Any]) -> int:
    try:
        return int(field.get("ordinal") or 0)
    except (TypeError, ValueError):
        return 0


def _localized_bool(value: Any, labels: dict[str, Any]) -> str:
    return labels["yes"] if bool(value) else labels["no"]


def _headers(labels: dict[str, Any]) -> list[str]:
    return [
        labels["ordinal"],
        labels["primary_key"],
        labels["field_code"],
        labels["field_description"],
        labels["data_type"],
        labels["length"],
        labels["nullable"],
        labels["default_value"],
        labels["comment"],
    ]


def _field_values(field: dict[str, Any], labels: dict[str, Any]) -> list[Any]:
    return [
        _ordinal(field),
        _localized_bool(field.get("is_primary_key"), labels),
        _text(field.get("code")),
        _text(field.get("name")),
        _text(field.get("data_type")),
        _text(field.get("length")),
        _localized_bool(field.get("nullable"), labels),
        _text(field.get("default_value")),
        _text(field.get("comment")),
    ]


def _xlsx_text(cell: Any, value: Any) -> None:
    cell.value = _text(value)
    # Treat user-provided PDM text as a literal, even if it begins with '='.
    cell.data_type = "s"


def _build_xlsx(detail: dict[str, Any], labels: dict[str, Any]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:
        worksheet = workbook.create_sheet()
    worksheet.title = labels["sheet"]
    worksheet.sheet_view.showGridLines = False

    _xlsx_text(worksheet["A1"], labels["table_name"])
    _xlsx_text(worksheet["B1"], detail.get("name"))
    _xlsx_text(worksheet["C1"], labels["table_code"])
    _xlsx_text(worksheet["D1"], detail.get("code"))
    _xlsx_text(worksheet["A2"], labels["table_description"])
    _xlsx_text(worksheet["B2"], detail.get("comment"))
    worksheet.merge_cells("B2:I2")

    metadata_fill = PatternFill("solid", fgColor="EAF2FF")
    for coordinate in ("A1", "C1", "A2"):
        metadata_cell = worksheet[coordinate]
        metadata_cell.fill = metadata_fill
        metadata_cell.font = Font(bold=True, color="315A91")
    for coordinate in ("B1", "D1", "B2"):
        worksheet[coordinate].alignment = Alignment(vertical="top", wrap_text=True)

    header_row = 4
    for column, header in enumerate(_headers(labels), start=1):
        header_cell = worksheet.cell(row=header_row, column=column)
        _xlsx_text(header_cell, header)
        header_cell.fill = PatternFill("solid", fgColor="3478E8")
        header_cell.font = Font(bold=True, color="FFFFFF")
        header_cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_number, field in enumerate(_fields(detail), start=header_row + 1):
        for column, value in enumerate(_field_values(field, labels), start=1):
            field_cell = cast(Cell, worksheet.cell(row=row_number, column=column))
            if isinstance(value, int):
                field_cell.value = value
            else:
                _xlsx_text(field_cell, value)
            field_cell.alignment = Alignment(vertical="top", wrap_text=column in {4, 9})

    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = f"A4:I{max(header_row, worksheet.max_row)}"
    for column, width in {
        "A": 9,
        "B": 12,
        "C": 26,
        "D": 32,
        "E": 18,
        "F": 14,
        "G": 12,
        "H": 20,
        "I": 48,
    }.items():
        worksheet.column_dimensions[column].width = width
    worksheet.row_dimensions[header_row].height = 24

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _protect_csv_formula(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    first_non_whitespace = value.lstrip(" \t\r\n")[:1]
    if first_non_whitespace in {"=", "+", "-", "@"}:
        return f"'{value}"
    return value


def _build_csv(detail: dict[str, Any], labels: dict[str, Any]) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(
        [
            labels["table_name"],
            labels["table_code"],
            labels["table_description"],
            *_headers(labels),
        ]
    )
    for field in _fields(detail):
        values = [
            _text(detail.get("name")),
            _text(detail.get("code")),
            _text(detail.get("comment")),
            *_field_values(field, labels),
        ]
        writer.writerow([_protect_csv_formula(value) for value in values])
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def _json_field(field: dict[str, Any]) -> dict[str, Any]:
    return {
        "ordinal": _ordinal(field),
        "is_primary_key": bool(field.get("is_primary_key")),
        "code": _text(field.get("code")),
        "description": _text(field.get("name")),
        "data_type": _text(field.get("data_type")),
        "length": _text(field.get("length")),
        "nullable": bool(field.get("nullable")),
        "default_value": _text(field.get("default_value")),
        "comment": _text(field.get("comment")),
    }


def _build_json(detail: dict[str, Any]) -> bytes:
    payload = {
        "table": {
            "name": _text(detail.get("name")),
            "code": _text(detail.get("code")),
            "description": _text(detail.get("comment")),
        },
        "fields": [_json_field(field) for field in _fields(detail)],
    }
    return (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def _markdown_cell(value: Any) -> str:
    return (
        _text(value)
        .replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "<br>")
        .replace("\n", "<br>")
        .replace("\r", "<br>")
    )


def _build_markdown(detail: dict[str, Any], labels: dict[str, Any]) -> bytes:
    table_name = _markdown_cell(detail.get("name") or detail.get("code"))
    table_code = _markdown_cell(detail.get("code"))
    description = _markdown_cell(detail.get("comment"))
    headers = _headers(labels)
    lines = [
        f"# {table_name}",
        "",
        f"- {labels['table_code']}: `{table_code}`",
        f"- {labels['table_description']}: {description}",
        "",
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for field in _fields(detail):
        values = [_markdown_cell(value) for value in _field_values(field, labels)]
        lines.append("| " + " | ".join(values) + " |")
    return ("\n".join(lines) + "\n").encode("utf-8")


def build_table_export(
    detail: dict[str, Any],
    export_format: TableExportFormat,
    language: Literal["zh-CN", "en-US"] = "zh-CN",
) -> tuple[bytes, str, str]:
    """Build an export containing table metadata and its ordered field structure."""
    labels = _LABELS[language]
    if export_format == "xlsx":
        return _build_xlsx(detail, labels), _XLSX_MEDIA_TYPE, ".xlsx"
    if export_format == "csv":
        return _build_csv(detail, labels), "text/csv; charset=utf-8", ".csv"
    if export_format == "json":
        return _build_json(detail), "application/json; charset=utf-8", ".json"
    if export_format == "md":
        return _build_markdown(detail, labels), "text/markdown; charset=utf-8", ".md"
    raise ValueError(f"Unsupported table export format: {export_format}")
